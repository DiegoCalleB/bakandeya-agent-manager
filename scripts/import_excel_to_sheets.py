import os
import sys
import uuid
import gspread
from openpyxl import load_workbook
from dotenv import load_dotenv

# Asegurar que el directorio raíz está en el path para las importaciones de lib
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

load_dotenv()

from lib.sheets import obtener_cliente_sheets, DOCUMENTO_SHEETS

# Columnas oficiales según AGENTS.md
COLUMNAS = [
    "id",
    "nombre_sala",
    "ciudad",
    "region",
    "aforo",
    "genero",
    "email_contacto",
    "fuente",
    "estado",
    "pitch_generado",
    "fecha_envio",
    "fecha_ultima_respuesta",
    "notas"
]

def map_aforo(val):
    """
    Convierte los tamaños del Excel ('S', 'M', 'L') a un aforo numérico aproximado.
    """
    if not val:
        return 0
    val_str = str(val).strip().upper()
    if val_str == 'S':
        return 150
    elif val_str == 'M':
        return 350
    elif val_str == 'L':
        return 800
    elif val_str == 'XL':
        return 1500
    elif 'S/M' in val_str:
        return 250
    elif 'M/L' in val_str:
        return 500
    try:
        # Si ya es un número, devolverlo como int
        return int(val)
    except ValueError:
        return 0

def leer_excel_leads():
    ruta = "Inputs/Bakandeya_Salas_Festivales.xlsx"
    if not os.path.exists(ruta):
        print(f"[ERROR] Archivo Excel no encontrado en: {ruta}")
        return []
        
    wb = load_workbook(ruta, data_only=True)
    leads_importados = []
    
    # 1. Procesar Salas España
    if 'Salas España' in wb.sheetnames:
        hoja = wb['Salas España']
        filas = list(hoja.iter_rows(values_only=True))
        # Cabecera: ('Ciudad', 'Nombre', 'Dirección', 'Teléfono', 'Rating Google', 'Nº reseñas', 'Tamaño aprox.', 'Encaje género (orientativo)', 'Verificado (sí/no)', 'Notas')
        for f in filas[1:]:
            if not f or not f[1]: # Saltear vacíos o sin nombre de sala
                continue
            ciudad, nombre, direccion, telefono, rating, resenas, tamano, genero, verificado, notas_excel = f[:10]
            
            aforo = map_aforo(tamano)
            notas_completas = (
                f"Dirección: {direccion or 'N/A'}. "
                f"Teléfono: {telefono or 'N/A'}. "
                f"Google Rating: {rating or 'N/A'} ({resenas or 0} reseñas). "
                f"Tamaño Excel: {tamano or 'N/A'}. "
                f"Notas Excel: {notas_excel or ''}"
            ).strip()
            
            leads_importados.append({
                "id": str(uuid.uuid4())[:8],
                "nombre_sala": nombre,
                "ciudad": ciudad,
                "region": "España",
                "aforo": aforo,
                "genero": genero or "Variado",
                "email_contacto": "", # El scout o humano lo buscará
                "fuente": "Excel: Salas España",
                "estado": "nuevo",
                "pitch_generado": "",
                "fecha_envio": "",
                "fecha_ultima_respuesta": "",
                "notas": notas_completas
            })
            
    # 2. Procesar Salas Europa cercana
    if 'Salas Europa cercana' in wb.sheetnames:
        hoja = wb['Salas Europa cercana']
        filas = list(hoja.iter_rows(values_only=True))
        # Cabecera: ('Ciudad', 'País', 'Nombre', 'Dirección', 'Teléfono', 'Rating Google', 'Nº reseñas', 'Tamaño aprox.', 'Encaje género (orientativo)', 'Verificado (sí/no)', 'Notas')
        for f in filas[1:]:
            if not f or not f[2]:
                continue
            ciudad, pais, nombre, direccion, telefono, rating, resenas, tamano, genero, verificado, notas_excel = f[:11]
            
            aforo = map_aforo(tamano)
            notas_completas = (
                f"Dirección: {direccion or 'N/A'}. "
                f"Teléfono: {telefono or 'N/A'}. "
                f"Google Rating: {rating or 'N/A'} ({resenas or 0} reseñas). "
                f"Tamaño Excel: {tamano or 'N/A'}. "
                f"Notas Excel: {notas_excel or ''}"
            ).strip()
            
            leads_importados.append({
                "id": str(uuid.uuid4())[:8],
                "nombre_sala": nombre,
                "ciudad": ciudad,
                "region": pais or "Europa",
                "aforo": aforo,
                "genero": genero or "Variado",
                "email_contacto": "",
                "fuente": f"Excel: Salas Europa ({pais})",
                "estado": "nuevo",
                "pitch_generado": "",
                "fecha_envio": "",
                "fecha_ultima_respuesta": "",
                "notas": notas_completas
            })

    # 3. Procesar Festivales y concursos
    if 'Festivales y concursos' in wb.sheetnames:
        hoja = wb['Festivales y concursos']
        filas = list(hoja.iter_rows(values_only=True))
        # Fila 1 es aviso, Fila 2 es cabecera: ('Nombre', 'Ubicación', 'Época aprox.', 'Encaje género', 'Tipo / tamaño', 'Notas - verificar antes de usar')
        for f in filas[2:]:
            if not f or not f[0] or f[0] == 'Nombre':
                continue
            nombre, ubicacion, epoca, genero, tamano, notas_excel = f[:6]
            
            notas_completas = (
                f"Época aproximada: {epoca or 'N/A'}. "
                f"Tamaño/Tipo: {tamano or 'N/A'}. "
                f"Notas Excel: {notas_excel or ''}"
            ).strip()
            
            leads_importados.append({
                "id": str(uuid.uuid4())[:8],
                "nombre_sala": nombre,
                "ciudad": ubicacion or "Varios",
                "region": "España / Europa",
                "aforo": 0, # Festivales suelen tener aforos no mapeados a salas
                "genero": genero or "Variado",
                "email_contacto": "",
                "fuente": "Excel: Festivales y concursos",
                "estado": "nuevo",
                "pitch_generado": "",
                "fecha_envio": "",
                "fecha_ultima_respuesta": "",
                "notas": notas_completas
            })
            
    return leads_importados

def importar_a_google_sheets():
    print("[INFO] Leyendo leads del archivo Excel local...")
    leads = leer_excel_leads()
    print(f"[INFO] Se parsearon {len(leads)} leads en total.")
    
    if not leads:
        print("[WARN] No hay leads para importar.")
        return
        
    try:
        print("[INFO] Conectando a Google Sheets API...")
        client = obtener_cliente_sheets()
        
        # Intentar abrir el documento
        try:
            doc = client.open(DOCUMENTO_SHEETS)
            print(f"[SUCCESS] Documento '{DOCUMENTO_SHEETS}' abierto correctamente.")
        except gspread.exceptions.SpreadsheetNotFound:
            print(f"[WARN] El documento '{DOCUMENTO_SHEETS}' no se encontró.")
            print("[INFO] Creando nueva hoja de cálculo...")
            doc = client.create(DOCUMENTO_SHEETS)
            print(f"[SUCCESS] Documento '{DOCUMENTO_SHEETS}' creado. ID: {doc.id}")
            print(
                f"[IMPORTANT] Como la cuenta de servicio creó este documento, DEBES compartirlo con tu correo personal.\n"
                f"Para ello, comparte el documento de Google Sheets con tu correo electrónico."
            )
            
        # Comprobar si existe la pestaña 'leads'
        worksheets_names = [w.title for w in doc.worksheets()]
        if 'leads' in worksheets_names:
            sheet = doc.worksheet('leads')
            print("[INFO] Pestaña 'leads' existente detectada. Se procederá a vaciarla e importar desde cero.")
            sheet.clear()
        else:
            sheet = doc.add_worksheet(title='leads', rows=1000, cols=len(COLUMNAS))
            print("[SUCCESS] Pestaña 'leads' creada con éxito.")
            # Borrar la pestaña por defecto si existe (Sheet1 o Hoja1)
            for w in doc.worksheets():
                if w.title in ['Sheet1', 'Hoja1']:
                    doc.del_worksheet(w)
                    break
        
        # Escribir cabecera
        sheet.append_row(COLUMNAS)
        
        # Preparar filas para inserción masiva
        filas_a_insertar = []
        for lead in leads:
            fila = [lead[col] for col in COLUMNAS]
            filas_a_insertar.append(fila)
            
        # Inserción masiva de leads
        sheet.append_rows(filas_a_insertar)
        print(f"[SUCCESS] Importación completada. Se insertaron {len(filas_a_insertar)} leads en la Google Sheet.")
        
    except Exception as e:
        print(f"[ERROR] Error al importar a Google Sheets: {e}")
        print("\nConsejo: Si el error es 403, asegúrate de haber habilitado la 'Google Drive API' en Google Cloud Console.")

if __name__ == "__main__":
    importar_a_google_sheets()

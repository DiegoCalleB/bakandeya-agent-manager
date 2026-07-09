import os
from openpyxl import load_workbook

def inspeccionar_excel():
    ruta = "Inputs/Bakandeya_Salas_Festivales.xlsx"
    if not os.path.exists(ruta):
        print(f"No se encontró el archivo en: {os.path.abspath(ruta)}")
        return
        
    wb = load_workbook(ruta, read_only=True)
    print(f"Hojas encontradas en el Excel: {wb.sheetnames}")
    
    for nombre_hoja in wb.sheetnames:
        hoja = wb[nombre_hoja]
        print(f"\n--- Hoja: {nombre_hoja} ---")
        filas = list(hoja.iter_rows(values_only=True))
        print(f"Total de filas: {len(filas)}")
        if filas:
            print(f"Cabecera (Fila 1): {filas[0]}")
            print("Primeras 3 filas de datos:")
            for idx, f in enumerate(filas[1:4], 1):
                print(f"  Fila {idx}: {f}")

if __name__ == "__main__":
    inspeccionar_excel()
